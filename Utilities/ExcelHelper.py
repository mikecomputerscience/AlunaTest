from openpyxl import *


def get_value(sheet, row, col):
    return sheet.cell(row, col).value


def get_dic(file_name, sheet_name):
    wb = load_workbook(file_name)
    ws = wb[sheet_name]
    ROW = ws.max_row
    COL = ws.max_column
    dic = {}
    for r in range(2, ROW + 1):
        the_id = get_value(ws, r, 1)
        dic[the_id] = {}
        for c in range(2, COL + 1):
            key = get_value(ws, 1, c)
            value = get_value(ws, r, c)
            dic[the_id][key] = value
    return dic
